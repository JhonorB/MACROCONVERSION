import os
from django.test import TestCase
import pandas as pd
import openpyxl
from api.services.procesador_sire import preparar_dataframes
from api.services.generador_excel import guardar_excel_final
from api.services.verificador_deudas import verificar_deudas
from django.conf import settings

class ProcesadorSireTests(TestCase):
    def setUp(self):
        self.tmp_dir = os.path.join(settings.BASE_DIR, 'tmp_tests')
        os.makedirs(self.tmp_dir, exist_ok=True)

    def tearDown(self):
        import shutil
        if os.path.exists(self.tmp_dir):
            shutil.rmtree(self.tmp_dir)

    def test_deteccion_tasa_igv_18_y_10_5(self):
        csv_content = (
            "RUC,Apellidos y Nombres o Razón social,Periodo,CAR SUNAT,Fecha de emisión,Fecha Vcto/Pago,Tipo CP/Doc.,Serie del CDP,Año,Nro CP o Doc. Nro Inicial (Rango),Nro Final (Rango),Tipo Doc Identidad,Nro Doc Identidad,Apellidos Nombres/ Razón  Social,BI Gravado DG,IGV / IPM DG,BI Gravado DGNG,IGV / IPM DGNG,BI Gravado DNG,IGV / IPM DNG,Valor Adq. NG,ISC,ICBPER,Otros Trib/ Cargos,Total CP,Moneda,Tipo de Cambio\n"
            "20543448568,EMPRESA DECLARANTE,202607,,01/07/2026,,01,F001,,101,,6,20111111111,PROVEEDOR TASA 18,100.00,18.00,0,0,0,0,0,0,0,0,118.00,PEN,1.000\n"
            "20543448568,EMPRESA DECLARANTE,202607,,02/07/2026,,01,F001,,102,,6,20222222222,PROVEEDOR TASA 10.5,100.00,10.50,0,0,0,0,0,0,0,0,110.50,PEN,1.000\n"
        )
        csv_file = os.path.join(self.tmp_dir, "compras_tasas.csv")
        with open(csv_file, 'w', encoding='utf-8') as f:
            f.write(csv_content)

        cuentas_asignadas = {
            "F001-101-20111111111": "6011001",
            "F001-102-20222222222": "6011001"
        }

        res = preparar_dataframes(
            csv_file,
            tipo="COM",
            codigo_empresa="0001",
            año="26",
            mes="07",
            cuentas_asignadas=cuentas_asignadas
        )
        df_forexel = res[0]

        self.assertEqual(len(df_forexel), 2)
        self.assertEqual(df_forexel.iloc[0]["TASIGV"], "18.00")
        self.assertEqual(df_forexel.iloc[1]["TASIGV"], "10.50")

    def test_asignacion_cencos_cuentas_gasto(self):
        csv_content = (
            "RUC,Apellidos y Nombres o Razón social,Periodo,CAR SUNAT,Fecha de emisión,Fecha Vcto/Pago,Tipo CP/Doc.,Serie del CDP,Año,Nro CP o Doc. Nro Inicial (Rango),Nro Final (Rango),Tipo Doc Identidad,Nro Doc Identidad,Apellidos Nombres/ Razón  Social,BI Gravado DG,IGV / IPM DG,BI Gravado DGNG,IGV / IPM DGNG,BI Gravado DNG,IGV / IPM DNG,Valor Adq. NG,ISC,ICBPER,Otros Trib/ Cargos,Total CP,Moneda,Tipo de Cambio\n"
            "20543448568,EMPRESA DECLARANTE,202607,,01/07/2026,,01,F001,,201,,6,20333333333,PROVEEDOR TRANSPORTE,500.00,90.00,0,0,0,0,0,0,0,0,590.00,PEN,1.000\n"
            "20543448568,EMPRESA DECLARANTE,202607,,02/07/2026,,01,F001,,202,,6,20444444444,PROVEEDOR MERCADERIA,1000.00,180.00,0,0,0,0,0,0,0,0,1180.00,PEN,1.000\n"
        )
        csv_file = os.path.join(self.tmp_dir, "compras_cencos.csv")
        with open(csv_file, 'w', encoding='utf-8') as f:
            f.write(csv_content)

        cuentas_asignadas = {
            "F001-201-20333333333": "6311001",
            "F001-202-20444444444": "6011001"
        }

        res = preparar_dataframes(
            csv_file,
            tipo="COM",
            codigo_empresa="0001",
            año="26",
            mes="07",
            cuentas_asignadas=cuentas_asignadas
        )
        df_forexel = res[0]

        self.assertEqual(df_forexel.iloc[0]["CTACOM"], "6311001")
        self.assertEqual(df_forexel.iloc[0]["CENCOS"], "V0000001")
        self.assertEqual(df_forexel.iloc[1]["CTACOM"], "6011001")
        self.assertNotEqual(df_forexel.iloc[1]["CENCOS"], "V0000001")

    def test_tratamiento_boletas_ventas(self):
        csv_content = (
            "RUC,Apellidos y Nombres o Razón social,Periodo,CAR SUNAT,Fecha de emisión,Fecha Vcto/Pago,Tipo CP/Doc.,Serie del CDP,Año,Nro CP o Doc. Nro Inicial (Rango),Nro Final (Rango),Tipo Doc Identidad,Nro Doc Identidad,Apellidos Nombres/ Razón  Social,BI Gravado DG,IGV / IPM DG,BI Gravado DGNG,IGV / IPM DGNG,BI Gravado DNG,IGV / IPM DNG,Valor Adq. NG,ISC,ICBPER,Otros Trib/ Cargos,Total CP,Moneda,Tipo de Cambio\n"
            "20543448568,EMPRESA DECLARANTE,202607,,01/07/2026,,03,B001,,501,,1,44556677,CLIENTE BOLETA,100.00,18.00,0,0,0,0,0,0,0,0,118.00,PEN,1.000\n"
        )
        csv_file = os.path.join(self.tmp_dir, "ventas_boletas.csv")
        with open(csv_file, 'w', encoding='utf-8') as f:
            f.write(csv_content)

        res = preparar_dataframes(
            csv_file,
            tipo="VTA",
            codigo_empresa="0001",
            año="26",
            mes="07",
            subtipo_vta="MERCADERIA"
        )
        df_forexel = res[0]
        row = df_forexel.iloc[0]

        self.assertEqual(row["TIPDOC"], "03")
        self.assertEqual(float(row["IGV"]), 0.00)
        self.assertEqual(float(row["VALVTA"]), 118.00)
        self.assertEqual(float(row["_BI_ORIG"]), 100.00)
        self.assertEqual(float(row["_IGV_ORIG"]), 18.00)
        self.assertEqual(float(row["_TOTAL_ORIG"]), 118.00)

    def test_lectura_csv_distintos_separadores_y_encodings(self):
        csv_semicolon = (
            "RUC;Apellidos y Nombres o Razón social;Periodo;CAR SUNAT;Fecha de emisión;Fecha Vcto/Pago;Tipo CP/Doc.;Serie del CDP;Año;Nro CP o Doc. Nro Inicial (Rango);Nro Final (Rango);Tipo Doc Identidad;Nro Doc Identidad;Apellidos Nombres/ Razón  Social;BI Gravado DG;IGV / IPM DG;BI Gravado DGNG;IGV / IPM DGNG;BI Gravado DNG;IGV / IPM DNG;Valor Adq. NG;ISC;ICBPER;Otros Trib/ Cargos;Total CP;Moneda;Tipo de Cambio\n"
            "20543448568;EMPRESA ÑANDÚ S.A.C.;202607;;01/07/2026;;01;E001;;801;;6;20555555555;PROVEEDOR CAÑÓN S.R.L.;200.00;36.00;0;0;0;0;0;0;0;0;236.00;PEN;1.000\n"
        )
        csv_file1 = os.path.join(self.tmp_dir, "semicolon_latin.csv")
        with open(csv_file1, 'wb') as f:
            f.write(csv_semicolon.encode("latin1"))

        res1 = preparar_dataframes(
            csv_file1,
            tipo="COM",
            codigo_empresa="0001",
            año="26",
            mes="07",
            cuentas_asignadas={"E001-801-20555555555": "6011001"}
        )
        df1 = res1[0]
        self.assertEqual(len(df1), 1)
        self.assertEqual(df1.iloc[0]["CODENT"], "20555555555")

        csv_pipe = (
            "RUC|Apellidos y Nombres o Razón social|Periodo|CAR SUNAT|Fecha de emisión|Fecha Vcto/Pago|Tipo CP/Doc.|Serie del CDP|Año|Nro CP o Doc. Nro Inicial (Rango)|Nro Final (Rango)|Tipo Doc Identidad|Nro Doc Identidad|Apellidos Nombres/ Razón  Social|BI Gravado DG|IGV / IPM DG|BI Gravado DGNG|IGV / IPM DGNG|BI Gravado DNG|IGV / IPM DNG|Valor Adq. NG|ISC|ICBPER|Otros Trib/ Cargos|Total CP|Moneda|Tipo de Cambio\n"
            "20543448568|EMPRESA DECLARANTE|202607||01/07/2026||01|F002||901||6|20666666666|PROVEEDOR PIPE S.A.|300.00|54.00|0|0|0|0|0|0|0|0|354.00|PEN|1.000\n"
        )
        csv_file2 = os.path.join(self.tmp_dir, "pipe_utf8.csv")
        with open(csv_file2, 'w', encoding='utf-8') as f:
            f.write(csv_pipe)

        res2 = preparar_dataframes(
            csv_file2,
            tipo="COM",
            codigo_empresa="0001",
            año="26",
            mes="07",
            cuentas_asignadas={"F002-901-20666666666": "6011001"}
        )
        df2 = res2[0]
        self.assertEqual(len(df2), 1)
        self.assertEqual(df2.iloc[0]["SERIE"], "F002")

    def test_ordenamiento_notas_credito(self):
        csv_content = (
            "RUC,Apellidos y Nombres o Razón social,Periodo,CAR SUNAT,Fecha de emisión,Fecha Vcto/Pago,Tipo CP/Doc.,Serie del CDP,Año,Nro CP o Doc. Nro Inicial (Rango),Nro Final (Rango),Tipo Doc Identidad,Nro Doc Identidad,Apellidos Nombres/ Razón  Social,BI Gravado DG,IGV / IPM DG,BI Gravado DGNG,IGV / IPM DGNG,BI Gravado DNG,IGV / IPM DNG,Valor Adq. NG,ISC,ICBPER,Otros Trib/ Cargos,Total CP,Moneda,Tipo de Cambio,Fecha Emisión Doc Modificado,Tipo CP Modificado,Serie CP Modificado,COD. DAM O DSI,Nro CP Modificado\n"
            "20543448568,EMPRESA DECLARANTE,202607,,05/07/2026,,07,FC01,,10,,6,20777777777,PROVEEDOR NC,50.00,9.00,0,0,0,0,0,0,0,0,59.00,PEN,1.000,01/07/2026,01,F001,,101\n"
            "20543448568,EMPRESA DECLARANTE,202607,,01/07/2026,,01,F001,,101,,6,20777777777,PROVEEDOR FACTURA,500.00,90.00,0,0,0,0,0,0,0,0,590.00,PEN,1.000,,,,,\n"
        )
        csv_file = os.path.join(self.tmp_dir, "compras_nc.csv")
        with open(csv_file, 'w', encoding='utf-8') as f:
            f.write(csv_content)

        cuentas_asignadas = {
            "FC01-10-20777777777": "6011001",
            "F001-101-20777777777": "6011001"
        }

        res = preparar_dataframes(
            csv_file,
            tipo="COM",
            codigo_empresa="0001",
            año="26",
            mes="07",
            cuentas_asignadas=cuentas_asignadas
        )
        df_forexel, df_docref, _, cantidad_nc, _, _, _ = res

        self.assertEqual(cantidad_nc, 1)
        ruta_generada = guardar_excel_final(df_forexel, df_docref, [], "COM", "compras_nc.xlsx")
        
        ruta_abs = os.path.join(settings.MEDIA_ROOT, ruta_generada)
        wb = openpyxl.load_workbook(ruta_abs)
        self.assertIn("FOREXEL", wb.sheetnames)
        self.assertIn("DOCREF", wb.sheetnames)

        sheet_forexel = wb["FOREXEL"]
        tipdoc_fila2 = str(sheet_forexel.cell(row=2, column=6).value)
        tipdoc_fila3 = str(sheet_forexel.cell(row=3, column=6).value)
        self.assertEqual(tipdoc_fila2, "01")
        self.assertEqual(tipdoc_fila3, "07")

    def test_extraccion_contraparte_proveedor_no_empresa(self):
        csv_content = (
            "RUC,Apellidos y Nombres o Razón social,Periodo,CAR SUNAT,Fecha de emisión,Fecha Vcto/Pago,Tipo CP/Doc.,Serie del CDP,Año,Nro CP o Doc. Nro Inicial (Rango),Nro Final (Rango),Tipo Doc Identidad,Nro Doc Identidad,Apellidos Nombres/ Razón  Social,BI Gravado DG,IGV / IPM DG,BI Gravado DGNG,IGV / IPM DGNG,BI Gravado DNG,IGV / IPM DNG,Valor Adq. NG,ISC,ICBPER,Otros Trib/ Cargos,Total CP,Moneda,Tipo de Cambio\n"
            "20543448568,LA CASA OXAPAMPINA A & K E.I.R.L.,202607,,01/07/2026,,01,E001,,10003,,6,20612744425,TRANSPORTE LOGISTICO YSABELLA S.A.C.,1429.19,257.25,0,0,0,0,0,0,0,0,1686.44,PEN,1.000\n"
        )
        csv_file = os.path.join(self.tmp_dir, "compras_oxapampina.csv")
        with open(csv_file, 'w', encoding='utf-8') as f:
            f.write(csv_content)

        res = preparar_dataframes(
            csv_file,
            tipo="COM",
            codigo_empresa="0001",
            año="26",
            mes="07",
            cuentas_asignadas=None
        )
        facturas = res[5]

        self.assertEqual(len(facturas), 1)
        self.assertEqual(facturas[0]["CODENT"], "20612744425")
        self.assertEqual(facturas[0]["RAZSOC"], "TRANSPORTE LOGISTICO YSABELLA S.A.C.")
        self.assertNotIn("OXAPAMPINA", facturas[0]["RAZSOC"])

    def test_verificar_deudas_sin_tokens(self):
        rucs = ["20111111111", "20222222222", "123"]
        resultado = verificar_deudas(rucs)
        self.assertIsInstance(resultado, list)

    def test_ventas_con_facturas_dolares_y_notas_credito(self):
        csv_content = (
            "Ruc,Razon Social,Periodo,CAR SUNAT,Fecha de emisión,Fecha Vcto/Pago,Tipo CP/Doc.,Serie del CDP,Nro CP o Doc. Nro Inicial (Rango),Nro Final (Rango),Tipo Doc Identidad,Nro Doc Identidad,Apellidos Nombres/ Razón Social,Valor Facturado Exportación,BI Gravada,Dscto BI,IGV / IPM,Dscto IGV / IPM,Mto Exonerado,Mto Inafecto,ISC,BI Grav IVAP,IVAP,ICBPER,Otros Tributos,Total CP,Moneda,Tipo Cambio,Fecha Emisión Doc Modificado,Tipo CP Modificado,Serie CP Modificado,Nro CP Modificado\n"
            "20520775839,J-SIMEC S.A.C.,202608,CAR1,03/08/2026,,01,E001,513,,6,10161254326,CLIENTE 1,0,1735.44,0,312.38,0,0,0,0,0,0,0,0,2047.82,PEN,1.000,,,,\n"
            "20520775839,J-SIMEC S.A.C.,202608,CAR2,03/08/2026,,07,E001,36,,6,10161254326,CLIENTE 1,0,-1735.44,0,-312.38,0,0,0,0,0,0,0,0,-2047.82,PEN,1.000,03/08/2026,01,E001,513\n"
            "20520775839,J-SIMEC S.A.C.,202608,CAR3,04/08/2026,,01,E001,515,,6,20547006101,CLIENTE USD,0,3643.54,0,655.84,0,0,0,0,0,0,0,0,4299.38,USD,3.402,,,,\n"
        )
        csv_file = os.path.join(self.tmp_dir, "ventas_usd_nc.csv")
        with open(csv_file, 'w', encoding='utf-8') as f:
            f.write(csv_content)

        res = preparar_dataframes(
            csv_file,
            tipo="VTA",
            codigo_empresa="0215",
            año="26",
            mes="08",
            subtipo_vta="MERCADERIA"
        )
        df_forexel = res[0]

        row_usd = df_forexel[df_forexel["NUMERO"] == "515"].iloc[0]
        self.assertEqual(float(row_usd["VALVTA"]), 1071.00)
        self.assertEqual(float(row_usd["IGV"]), 192.78)
        self.assertEqual(row_usd["MND"], "D")
        self.assertEqual(float(row_usd["_BI_ORIG"]), 3643.54)
        self.assertEqual(float(row_usd["_IGV_ORIG"]), 655.84)
